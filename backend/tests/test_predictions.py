"""押题端点测试：校验、隔离、大模型错误路径与成功路径（桩实现）。"""

from __future__ import annotations

import io
import uuid

import docx
import pytest

from app.api.deps import get_storage
from app.api.v1.endpoints.predictions import get_llm
from app.services.llm import (
    Evidence,
    LLMClient,
    PredictedQuestionDraft,
    PredictionDraft,
)
from app.services.storage import LocalStorage

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MD_MIME = "text/markdown"


class FakeLLM(LLMClient):
    """测试桩：记录调用并返回固定草稿。"""

    def __init__(
        self,
        draft: PredictionDraft | None = None,
        error: Exception | None = None,
    ) -> None:
        self.draft = draft
        self.error = error
        self.calls: list[dict] = []

    def generate_prediction(self, **kwargs) -> PredictionDraft:
        self.calls.append(kwargs)
        if self.error is not None:
            raise self.error
        assert self.draft is not None
        return self.draft


@pytest.fixture
def uploads_dir(tmp_path) -> LocalStorage:
    local = LocalStorage(tmp_path / "uploads")
    from app.main import app

    app.dependency_overrides[get_storage] = lambda: local
    yield local
    from app.main import app

    app.dependency_overrides.pop(get_storage, None)


@pytest.fixture
def fake_llm() -> FakeLLM:
    evidence = [
        Evidence(source="resume", ref="我的简历.docx", excerpt="五年 Python 后端经验"),
        Evidence(source="job_description", ref="岗位描述", excerpt="熟悉 FastAPI"),
        Evidence(source="knowledge", ref="tech.md", excerpt="FastAPI 与数据库设计"),
    ]
    draft = PredictionDraft(
        questions=[
            PredictedQuestionDraft(
                question=f"请介绍你在 FastAPI 项目中的经验（{index}）",
                difficulty="medium",
                reference_answer="要点：项目背景、承担职责、遇到的问题与解决思路。",
                evidence=evidence,
            )
            for index in range(2)
        ]
    )
    fake = FakeLLM(draft=draft)
    from app.main import app

    app.dependency_overrides[get_llm] = lambda: fake
    yield fake
    from app.main import app

    app.dependency_overrides.pop(get_llm, None)


def _docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("张三")
    document.add_paragraph("五年 Python 后端经验，熟悉 FastAPI")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _login(client, user: dict) -> str:
    response = client.post(
        "/api/v1/auth/login",
        json={"email": user["email"], "password": user["password"]},
    )
    assert response.status_code == 200, response.text
    token = response.json()["access_token"]
    client.headers["Authorization"] = f"Bearer {token}"
    return token


def _upload_resume(client) -> dict:
    response = client.post(
        "/api/v1/resumes", files={"file": ("我的简历.docx", _docx_bytes(), DOCX_MIME)}
    )
    assert response.status_code == 201, response.text
    return response.json()


def _payload(resume_id: str) -> dict:
    return {
        "resume_id": resume_id,
        "target_role": "Python 后端工程师",
        "years_of_experience": "3-5年",
        "job_description": "负责后端服务开发，使用 Python 与 FastAPI，熟悉数据库设计。",
        "programming_language": "Python",
    }


class TestCreatePrediction:
    def test_requires_auth(self, client) -> None:
        response = client.post("/api/v1/predictions", json=_payload(str(uuid.uuid4())))
        assert response.status_code == 401

    def test_missing_required_field_reports_field_errors(
        self, client, registered_user, uploads_dir, fake_llm
    ) -> None:
        _login(client, registered_user)
        payload = _payload(str(uuid.uuid4()))
        del payload["target_role"]
        response = client.post("/api/v1/predictions", json=payload)
        assert response.status_code == 422
        assert response.json()["code"] == "VALIDATION_ERROR"
        fields = [item["field"] for item in response.json()["details"]]
        assert "target_role" in fields

    def test_job_description_over_limit_rejected(
        self, client, registered_user, uploads_dir, fake_llm
    ) -> None:
        _login(client, registered_user)
        payload = _payload(str(uuid.uuid4()))
        payload["job_description"] = "描" * 2001
        response = client.post("/api/v1/predictions", json=payload)
        assert response.status_code == 422

    def test_resume_not_found(self, client, registered_user, uploads_dir, fake_llm) -> None:
        _login(client, registered_user)
        response = client.post("/api/v1/predictions", json=_payload(str(uuid.uuid4())))
        assert response.status_code == 404

    def test_unparsed_resume_rejected(
        self, client, registered_user, uploads_dir, fake_llm
    ) -> None:
        _login(client, registered_user)
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        buffer = io.BytesIO()
        writer.write(buffer)
        blank = client.post(
            "/api/v1/resumes", files={"file": ("blank.pdf", buffer.getvalue(), "application/pdf")}
        ).json()

        response = client.post("/api/v1/predictions", json=_payload(blank["id"]))
        assert response.status_code == 409
        assert response.json()["code"] == "RESUME_NOT_PARSED"

    def test_create_success_with_knowledge_reference(
        self, client, registered_user, uploads_dir, fake_llm, db_session
    ) -> None:
        _login(client, registered_user)
        resume = _upload_resume(client)

        knowledge = client.post(
            "/api/v1/knowledge/documents",
            files={
                "file": (
                    "tech.md",
                    "后端工程师面试要点：深入理解 Python，熟悉 FastAPI 框架与数据库设计。".encode(),
                    MD_MIME,
                )
            },
            data={"is_public": "true"},
        ).json()

        response = client.post("/api/v1/predictions", json=_payload(resume["id"]))
        assert response.status_code == 201, response.text
        body = response.json()
        assert body["status"] == "completed"
        assert body["model_name"]
        assert body["prompt_version"]
        assert body["generated_at"]
        assert len(body["questions"]) == 2
        assert body["questions"][0]["position"] == 1

        # 知识库命中依据带文档标识
        knowledge_evidence = [
            ev
            for question in body["questions"]
            for ev in question["evidence"]
            if ev["source"] == "knowledge"
        ]
        assert knowledge_evidence
        assert knowledge_evidence[0]["ref_id"] == knowledge["id"]

        # 大模型调用入参包含简历文本与知识库片段
        call = fake_llm.calls[0]
        assert "五年 Python 后端经验" in call["resume_text"]
        assert call["knowledge_snippets"]
        assert call["question_count"] == 8

    def test_empty_model_output_rejected(
        self, client, registered_user, uploads_dir
    ) -> None:
        empty_llm = FakeLLM(draft=PredictionDraft(questions=[]))
        from app.main import app

        app.dependency_overrides[get_llm] = lambda: empty_llm
        try:
            _login(client, registered_user)
            resume = _upload_resume(client)
            response = client.post("/api/v1/predictions", json=_payload(resume["id"]))
            assert response.status_code == 502
            assert response.json()["code"] == "MODEL_OUTPUT_INVALID"
        finally:
            app.dependency_overrides.pop(get_llm, None)

    def test_llm_runtime_error_becomes_unavailable(
        self, client, registered_user, uploads_dir
    ) -> None:
        broken_llm = FakeLLM(draft=None, error=RuntimeError("connection reset"))
        from app.main import app

        app.dependency_overrides[get_llm] = lambda: broken_llm
        try:
            _login(client, registered_user)
            resume = _upload_resume(client)
            response = client.post("/api/v1/predictions", json=_payload(resume["id"]))
            assert response.status_code == 503
            assert response.json()["code"] == "MODEL_UNAVAILABLE"
        finally:
            app.dependency_overrides.pop(get_llm, None)

    def test_missing_api_key_reports_unavailable(
        self, client, registered_user, uploads_dir, monkeypatch
    ) -> None:
        # 不注入桩：走真实 OpenAICompatLLM，无 Key 时返回 503
        monkeypatch.delenv("LLM_API_KEY", raising=False)
        _login(client, registered_user)
        resume = _upload_resume(client)
        response = client.post("/api/v1/predictions", json=_payload(resume["id"]))
        assert response.status_code == 503
        assert response.json()["code"] == "MODEL_UNAVAILABLE"


class TestPredictionHistory:
    def test_history_visible_to_owner_only(
        self, client, db_session, registered_user, uploads_dir, fake_llm
    ) -> None:
        _login(client, registered_user)
        resume = _upload_resume(client)
        created = client.post("/api/v1/predictions", json=_payload(resume["id"])).json()

        listing = client.get("/api/v1/predictions")
        assert listing.status_code == 200
        assert len(listing.json()) == 1
        assert listing.json()[0]["question_count"] == 2

        detail = client.get(f"/api/v1/predictions/{created['id']}")
        assert detail.status_code == 200
        assert detail.json()["job_description"].startswith("负责后端服务开发")

        # 另一个用户访问 → 404，不泄露存在性
        other_email = f"other-{uuid.uuid4().hex[:6]}@example.com"
        client.post(
            "/api/v1/auth/register", json={"email": other_email, "password": "Passw0rd1"}
        )
        _login(client, {"email": other_email, "password": "Passw0rd1"})
        assert client.get(f"/api/v1/predictions/{created['id']}").status_code == 404
        assert client.get("/api/v1/predictions").json() == []
